"""토지이음 공개자료 → 땅의 가치용 파일 생성 (대구·경북)
사용: python scripts/build_land_rules.py <행위제한정보.xlsx> <개발행위허가정보.csv>
출력: docs/landrules.json  (시군별 조례 + 법령 기준, 관심 건축물만)
      docs/permits/<시군코드5>.json  (리·동별 최근 개발행위허가 집계와 사례)
자료: 토지이음 데이터개방 '토지이용규제 행위제한정보', '개발행위허가정보' (국토교통부)"""
import sys, os, re, json, csv, collections
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
ZONES = ["제1종전용주거지역", "제2종전용주거지역", "제1종일반주거지역", "제2종일반주거지역", "제3종일반주거지역", "준주거지역",
         "중심상업지역", "일반상업지역", "근린상업지역", "유통상업지역", "전용공업지역", "일반공업지역", "준공업지역",
         "보전녹지지역", "생산녹지지역", "자연녹지지역", "보전관리지역", "생산관리지역", "계획관리지역", "농림지역", "자연환경보전지역", "관리지역"]
# 관심 건축물: (키, 토지이용명, 별표1 호목 일부)
ITEMS = [("house", "단독주택。", "제1호가목"), ("warehouse", "창고。", "제18호가목"), ("retail", "일용품을 판매하는 소매점。", "제3호가목"),
         ("cafe", "휴게음식점。", "제3호나목"), ("restaurant", "일반음식점。", "제4호자목"), ("factory", "공장。", "제17호")]
FLOOR = re.compile(r"(단,\s*)?4층(을)? 초과하는[^.]*?금지한다\.|4층 이하의 범위[^.]*?(금지한다|한한다)\.|4층 이하의 건축물에 한한다\.|다만, 4층 이하의 범위[^.]*?한한다\.")


def norm(v):
    return str(v).strip() if v is not None else ""


def classify(status, cond):
    """o 가능 / m 조건부 / x 금지. 층수 제한 문구는 조건으로 보지 않음"""
    c = FLOOR.sub("", cond).strip()
    if status.startswith("건축 가능") or status == "건축":
        if not c:
            return "o", ""
        return ("m" if re.search(r"한한다|한정|에 한함|경우에|만 해당", c) else "o"), c
    if status.startswith("건축 금지"):
        if not c:
            return "x", ""
        if re.search(r"다만|제외한다|지역에서는 금지|설치하는 것에 한한다|않은 지역", c):
            return "m", c
        return "x", c
    return "m", c


def pick(rows):
    """행 목록에서 관심 건축물만 {zone:{key:[cls, status, cond]}}"""
    out = collections.defaultdict(dict)
    for r in rows:
        zone = r[2]
        if zone not in ZONES:
            continue
        for key, name, mok in ITEMS:
            if r[4] == name and mok in r[3].replace(" ", ""):
                cls, c = classify(r[5], r[6])
                prev = out[zone].get(key)
                # 같은 항목이 여러 줄이면 더 엄격한 쪽(금지>조건부>가능)을 남김
                rank = {"x": 2, "m": 1, "o": 0}
                if not prev or rank[cls] > rank[prev[0]]:
                    out[zone][key] = [cls, r[5], c[:400]]
    return out


def build_rules(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    law = [[norm(x) for x in r[:7]] for r in wb["법령"].iter_rows(min_row=2, values_only=True)]
    res = {"source": os.path.basename(xlsx), "law": pick(law), "sgg": {}}
    for sh in ("조례_경북", "조례_대구"):
        rows = [[norm(x) for x in r[:7]] for r in wb[sh].iter_rows(min_row=2, values_only=True)]
        by = collections.defaultdict(list)
        for r in rows:
            by[r[0]].append(r)
        for code, rs in by.items():
            if len(rs) < 200:      # 일부 항목만 있는 코드(구 단위 소수 행)는 건너뜀
                continue
            res["sgg"][code] = {"name": rs[0][1], "rules": pick(rs)}
    return res


CATS = [("solar", r"태양광|발전시설|발전소"), ("house", r"단독주택|농가주택|주택"), ("warehouse", r"창고"),
        ("shop", r"근린생활|음식점|카페|소매점|휴게"), ("factory", r"공장|제조")]


def build_permits(csvf, since="2014"):
    agg = collections.defaultdict(lambda: {"n": collections.Counter(), "ex": collections.defaultdict(list)})
    with open(csvf, encoding="cp949", errors="replace", newline="") as f:
        rd = csv.reader(f)
        next(rd)
        for row in rd:
            pnu = row[0].strip()
            if pnu[:2] not in ("27", "47") or len(pnu) != 19:
                continue
            d = (row[13] or "").strip()
            if not d[:4].isdigit() or d[:4] < since:
                continue
            purpose = re.sub(r"\s+", " ", row[14]).strip()
            cat = next((k for k, rx in CATS if re.search(rx, purpose)), None)
            if not cat:
                continue
            a = agg[pnu[:10]]
            a["n"][cat] += 1
            loc = re.sub(r"\s+", " ", row[1]).strip()
            a["ex"][cat].append([d, loc.split(" ")[-1], purpose[:40], (row[4] or "").strip()])
    files = collections.defaultdict(dict)
    for li, a in agg.items():
        ex = {k: sorted(v, reverse=True)[:3] for k, v in a["ex"].items()}
        files[li[:5]][li] = {"n": dict(a["n"]), "ex": ex}
    return files, since


if __name__ == "__main__":
    xlsx, csvf = sys.argv[1], sys.argv[2]
    rules = build_rules(xlsx)
    with open(os.path.join(DOCS, "landrules.json"), "w", encoding="utf-8") as f:
        json.dump(rules, f, ensure_ascii=False, separators=(",", ":"))
    print("rules: sgg", len(rules["sgg"]), "law zones", len(rules["law"]))
    files, since = build_permits(csvf)
    os.makedirs(os.path.join(DOCS, "permits"), exist_ok=True)
    tot = 0
    for sgg, d in files.items():
        with open(os.path.join(DOCS, "permits", sgg + ".json"), "w", encoding="utf-8") as f:
            json.dump({"since": since, "source": os.path.basename(csvf), "li": d}, f, ensure_ascii=False, separators=(",", ":"))
        tot += len(d)
    print("permits: sgg files", len(files), "li", tot)
